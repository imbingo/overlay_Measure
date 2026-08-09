from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Dict, List, Optional

import pandas as pd
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .measurement_units import ellipse_metrics_um, mean_pixel_size_um, rotated_rect_size_um, scalar_px_to_um
from .geometry_models import GeometryRunResult
from .models import DetectionResult, MeasurementConfig, OverlayResult
from .quality_profiles import quality_profile_display


DETAIL_COLUMNS = {
    "roi_id": "ROI稳定ID",
    "roi_label": "ROI编号",
    "timestamp": "测量时间",
    "run_index": "测量次数",
    "measurement_mode": "测量模式",
    "upper_file": "上层/单图文件",
    "lower_file": "下层文件",
    "pixel_size_x_um": "像素尺寸X(μm/px)",
    "pixel_size_y_um": "像素尺寸Y(μm/px)",
    "registration_offset_x_um": "配准偏移X(μm)",
    "registration_offset_y_um": "配准偏移Y(μm)",
    "mark_id": "标记",
    "layer": "层",
    "center_x_um": "中心X(μm)",
    "center_y_um": "中心Y(μm)",
    "diameter_um": "直径/尺寸(μm)",
    "average_diameter_um": "平均直径(μm)",
    "maximum_diameter_um": "最大直径(μm)",
    "minimum_diameter_um": "最小直径(μm)",
    "diameter_pv_um": "直径PV(μm)",
    "diameter_mode": "直径定义",
    "fit_residual_um": "参考残差(μm)",
    "edge_point_count": "边缘点数",
    "confidence": "置信度",
    "fitting_mode": "拟合模式",
    "algorithm_path": "算法路径",
    "measurement_stage": "测量阶段",
    "quality_status": "质量状态",
    "quality_profile": "质量门槛",
    "quality_grade": "实际质量",
    "quality_details": "质量详情",
    "coverage": "覆盖率",
    "angular_coverage": "角度覆盖率",
    "maximum_gap_deg": "最大缺口角度(°)",
    "rejected_count": "剔除点数",
    "rejected_ratio": "异常点比例",
    "max_deviation_um": "最大轮廓偏差(μm)",
    "failure_reason": "失效原因",
    "recipe_validation_status": "配方验证状态",
    "detection_warning": "识别提示",
    "shape_width_um": "宽度(μm)",
    "shape_height_um": "高度(μm)",
    "shape_major_um": "长轴(μm)",
    "shape_minor_um": "短轴(μm)",
    "ellipse_diameter_um": "椭圆直径(μm)",
    "ellipse_roundness_um": "椭圆圆度(μm)",
    "shape_angle_deg": "角度(°)",
    "shape_aspect_ratio": "宽高比",
    "roi_type": "ROI类型",
    "roi_inner_ratio": "内环比例",
    "roi_inner_radius_um": "ROI内半径(μm)",
    "roi_outer_radius_um": "ROI外半径(μm)",
    "caliper_count": "卡尺数量",
    "caliper_width_um": "卡尺宽度(μm)",
    "search_direction": "搜索方向",
    "roi_target_edge": "边缘选择",
    "roi_angle_deg": "ROI角度(°)",
    "delta_x_um": "ΔX(μm)",
    "delta_y_um": "ΔY(μm)",
    "overlay_r_um": "Dxy/R(μm)",
    "result": "结果",
    "overlay_warning": "对位提示",
}


def _mode_cn(mode: str) -> str:
    return "双图模式" if mode == "Dual Image" else "单图模式"


def _layer_cn(layer: str) -> str:
    return {"upper": "上层", "lower": "下层"}.get(layer, layer)


def _result_cn(result: str) -> str:
    return {
        "Pass": "通过",
        "Fail": "超限",
        "Invalid": "无效",
        "Error": "异常",
        "Trial": "试测/不判定",
    }.get(result, result)


def _fit_cn(mode: str) -> str:
    return {
        "EdgeCenter": "边缘中心",
        "RegionCenter": "区域中心",
        "CaliperCircle": "卡尺找圆",
        "AutoCircle": "自动圆轮廓",
        "AutoRectangle": "自动方形轮廓",
        "ProductionCircle": "正式卡尺圆拟合",
        "ProductionRectangle": "正式四边卡尺拟合",
        "Auto": "自动",
        "Circle": "圆",
        "Ellipse": "椭圆",
        "Rectangle": "矩形/方孔",
    }.get(mode, mode)


def _roi_cn(value: str) -> str:
    return {
        "Annulus": "圆环",
        "Caliper Circle": "卡尺圆",
        "Rectangular Ring": "矩形环",
        "Circle": "圆",
        "Rectangle": "矩形",
        "Auto Full Image": "全图自动识别",
        "Auto Caliper Circle": "自动卡尺圆",
        "Auto Four-Side Caliper": "自动四边卡尺",
        "All Edges": "全部边缘",
        "Near Inner Boundary": "靠近内环",
        "Near Outer Boundary": "靠近外环",
        "Strongest Edge": "最强边缘",
    }.get(value, value)


def _direction_cn(value: str) -> str:
    return {
        "Inner to Outer": "由内向外",
        "Outer to Inner": "由外向内",
    }.get(value, value)


def build_detection_rows(
    detections: Dict[str, Dict[str, DetectionResult]],
    overlays: Dict[str, OverlayResult],
    config: MeasurementConfig,
    upper_file: str = "",
    lower_file: str = "",
    run_index: Optional[int] = None,
) -> List[dict]:
    rows = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    mean_scale_um = mean_pixel_size_um(config)
    for mark_id, layer_map in detections.items():
        overlay = overlays.get(mark_id)
        for detection_key, det in layer_map.items():
            width_px = det.shape_params.get("width_px")
            height_px = det.shape_params.get("height_px")
            angle_deg = float(det.shape_params.get("angle_deg", 0.0))
            width_um = height_um = None
            if width_px is not None and height_px is not None:
                width_um, height_um = rotated_rect_size_um(float(width_px), float(height_px), angle_deg, config)
            major_px = det.shape_params.get("major_px")
            minor_px = det.shape_params.get("minor_px")
            ellipse_metrics = ellipse_metrics_um(det.shape_params, config) if det.fitting_mode == "Ellipse" else {}
            row = {
                "timestamp": now,
                "run_index": run_index,
                "measurement_mode": _mode_cn(config.mode),
                "upper_file": upper_file,
                "lower_file": lower_file,
                "pixel_size_x_um": config.pixel_size_x_um,
                "pixel_size_y_um": config.pixel_size_y_um,
                "registration_offset_x_um": config.registration_offset_x_um,
                "registration_offset_y_um": config.registration_offset_y_um,
                "mark_id": mark_id,
                "roi_id": det.shape_params.get("roi_id", detection_key),
                "roi_label": det.shape_params.get("roi_label", ""),
                "layer": _layer_cn(det.layer),
                "center_x_um": det.center_x_um,
                "center_y_um": det.center_y_um,
                "diameter_um": det.diameter_um,
                "average_diameter_um": det.shape_params.get("average_diameter_um"),
                "maximum_diameter_um": det.shape_params.get("maximum_diameter_um"),
                "minimum_diameter_um": det.shape_params.get("minimum_diameter_um"),
                "diameter_pv_um": det.shape_params.get("diameter_pv_um"),
                "diameter_mode": {"Average": "平均直径", "Maximum": "最大直径"}.get(
                    det.shape_params.get("diameter_mode"),
                    det.shape_params.get("diameter_mode", ""),
                ),
                "fit_residual_um": det.residual_um,
                "edge_point_count": det.edge_point_count,
                "confidence": det.confidence,
                "fitting_mode": _fit_cn(det.fitting_mode),
                "algorithm_path": det.shape_params.get("algorithm_path", ""),
                "measurement_stage": "正式精测" if det.shape_params.get("measurement_stage") == "production_measurement" else "候选检测",
                "quality_status": {"Valid": "有效", "Invalid": "无效"}.get(det.shape_params.get("quality_status"), det.shape_params.get("quality_status", "")),
                "quality_profile": det.shape_params.get("quality_profile_label", ""),
                "quality_grade": det.shape_params.get("quality_grade", ""),
                "quality_details": det.shape_params.get("quality_details", ""),
                "coverage": det.shape_params.get("coverage"),
                "angular_coverage": det.shape_params.get("angular_coverage"),
                "maximum_gap_deg": det.shape_params.get("maximum_gap_deg"),
                "rejected_count": det.shape_params.get("rejected_count"),
                "rejected_ratio": det.shape_params.get("rejected_ratio"),
                "max_deviation_um": det.shape_params.get("max_deviation_um"),
                "failure_reason": det.shape_params.get("failure_reason"),
                "recipe_validation_status": {"Draft": "草稿/未验证", "Validated": "已验证/正式生产"}.get(det.shape_params.get("recipe_validation_status"), det.shape_params.get("recipe_validation_status", "")),
                "detection_warning": det.warning,
                "shape_width_um": width_um,
                "shape_height_um": height_um,
                "shape_major_um": det.ellipse_major_um if det.ellipse_major_um is not None else ellipse_metrics.get("ellipse_major_um"),
                "shape_minor_um": det.ellipse_minor_um if det.ellipse_minor_um is not None else ellipse_metrics.get("ellipse_minor_um"),
                "ellipse_diameter_um": det.ellipse_diameter_um if det.ellipse_diameter_um is not None else ellipse_metrics.get("ellipse_diameter_um"),
                "ellipse_roundness_um": det.ellipse_roundness_um if det.ellipse_roundness_um is not None else ellipse_metrics.get("ellipse_roundness_um"),
                "shape_angle_deg": det.shape_params.get("angle_deg"),
                "shape_aspect_ratio": det.shape_params.get("aspect_ratio"),
                "roi_type": _roi_cn(det.shape_params.get("roi_type")),
                "roi_inner_ratio": det.shape_params.get("roi_inner_ratio"),
                "roi_inner_radius_um": scalar_px_to_um(det.shape_params.get("roi_inner_radius_px"), config) if det.shape_params.get("roi_inner_radius_px") is not None else None,
                "roi_outer_radius_um": scalar_px_to_um(det.shape_params.get("roi_outer_radius_px"), config) if det.shape_params.get("roi_outer_radius_px") is not None else None,
                "caliper_count": det.shape_params.get("caliper_count"),
                "caliper_width_um": scalar_px_to_um(det.shape_params.get("caliper_width_px"), config) if det.shape_params.get("caliper_width_px") is not None else None,
                "search_direction": _direction_cn(det.shape_params.get("search_direction")),
                "roi_target_edge": _roi_cn(det.shape_params.get("roi_target_edge")),
                "roi_angle_deg": det.shape_params.get("roi_angle_deg"),
                "delta_x_um": overlay.delta_x_um if overlay else None,
                "delta_y_um": overlay.delta_y_um if overlay else None,
                "overlay_r_um": overlay.overlay_r_um if overlay else None,
                "result": _result_cn(overlay.result) if overlay else "",
                "overlay_warning": overlay.warning if overlay else "",
            }
            rows.append(row)
    return rows


def build_detection_failure_row(
    config: MeasurementConfig,
    run_index: int,
    mark_id: str,
    upper_file: str,
    lower_file: str,
    error: str,
    *,
    layer: str = "",
    roi_id: str = "",
    roi_index: Optional[int] = None,
    roi_label: str = "",
    status: str = "Error",
) -> dict:
    return {
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "run_index": run_index,
        "measurement_mode": _mode_cn(config.mode),
        "upper_file": upper_file,
        "lower_file": lower_file,
        "mark_id": mark_id,
        "roi_id": roi_id,
        "roi_label": roi_label or (f"ROI {roi_index}" if roi_index else ""),
        "layer": _layer_cn(layer),
        "quality_status": "无效" if status == "Invalid" else "异常",
        "failure_reason": error,
        "detection_warning": error,
        "result": "无效" if status == "Invalid" else "异常",
    }


def build_geometry_rows(result: GeometryRunResult, config: MeasurementConfig, run_index: Optional[int] = None) -> List[dict]:
    rows: List[dict] = []
    mean_scale = 0.5 * (config.pixel_size_x_um + config.pixel_size_y_um)
    for feature in result.features.values():
        rows.append({
            "测量次数": run_index,
            "类别": "要素",
            "编号": feature.feature_id,
            "名称": feature.name,
            "层": _layer_cn(feature.layer),
            "类型": feature.feature_type,
            "数值": None if feature.radius_px is None else 2.0 * feature.radius_px * mean_scale,
            "单位": "μm" if feature.radius_px is not None else "",
            "中心X(px)": None if feature.center_px is None else feature.center_px[0],
            "中心Y(px)": None if feature.center_px is None else feature.center_px[1],
            "状态": "有效" if feature.status == "Valid" else "无效",
            "质量": feature.quality,
            "算法路径": feature.algorithm_path,
            "提示": feature.error,
        })
    for coordinate in result.coordinate_systems.values():
        rows.append({
            "测量次数": run_index, "类别": "坐标系", "编号": coordinate.coordinate_id,
            "名称": coordinate.name, "层": _layer_cn(coordinate.layer), "类型": "自定义坐标系",
            "数值": coordinate.rotation_deg, "单位": "°", "状态": "有效" if coordinate.status == "Valid" else "无效",
            "质量": "", "算法路径": "物理坐标标定 → 自定义原点与轴", "提示": coordinate.error,
        })
    for measurement in result.measurements.values():
        rows.append({
            "测量次数": run_index, "类别": "测量", "编号": measurement.measurement_id,
            "名称": measurement.name, "层": _layer_cn(measurement.layer), "类型": measurement.measurement_type,
            "数值": measurement.value, "单位": measurement.unit,
            "状态": "有效" if measurement.status == "Valid" else "无效", "质量": measurement.quality,
            "算法路径": measurement.algorithm_path, "提示": measurement.error,
        })
    for label in result.coordinate_labels.values():
        rows.append({
            "测量次数": run_index, "类别": "坐标标注", "编号": label.label_id,
            "名称": label.name, "层": _layer_cn(label.layer), "类型": "圆心/点坐标",
            "数值": None, "单位": "μm", "坐标X(μm)": label.x_um, "坐标Y(μm)": label.y_um,
            "状态": "有效" if label.status == "Valid" else "无效", "质量": "",
            "算法路径": "要素中心 → 自定义坐标系", "提示": label.error,
        })
    return rows


def _autosize(ws):
    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 8
        for cell in column_cells:
            if cell.value is not None:
                max_len = max(max_len, min(36, len(str(cell.value)) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len


def _style_sheet(ws, fail_columns: Optional[List[str]] = None):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    fail_font = Font(color="9C0006", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    fail_columns = fail_columns or []
    headers = [cell.value for cell in ws[1]]
    fail_indexes = [headers.index(col) + 1 for col in fail_columns if col in headers]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
        for idx in fail_indexes:
            value = row[idx - 1].value
            if value in {"失败", "异常", "无效", "NG"} or (isinstance(value, str) and "超限" in value):
                row[idx - 1].fill = fail_fill
                row[idx - 1].font = fail_font
    _autosize(ws)


def resize_dimensions_preserving_aspect(width: float, height: float, max_width: float, max_height: float) -> tuple[int, int]:
    scale = min(max_width / max(width, 1.0), max_height / max(height, 1.0), 1.0)
    return int(round(width * scale)), int(round(height * scale))


def export_results(
    path: str,
    rows: List[dict],
    config: Optional[MeasurementConfig] = None,
    summary_rows: Optional[List[dict]] = None,
    mark_images: Optional[List[dict]] = None,
    repeatability_rows: Optional[List[dict]] = None,
    traceability_info: Optional[dict] = None,
    geometry_rows: Optional[List[dict]] = None,
) -> None:
    detail_df = pd.DataFrame(rows).rename(columns=DETAIL_COLUMNS)
    summary_df = pd.DataFrame(summary_rows or [])
    repeatability_df = pd.DataFrame(repeatability_rows or [])
    geometry_df = pd.DataFrame(geometry_rows or [])
    # Keep exported measurement results concise for production review.
    # Internal calculation remains full precision; only output tables are rounded.
    detail_df = detail_df.round(3)
    summary_df = summary_df.round(3)
    repeatability_df = repeatability_df.round(3)
    geometry_df = geometry_df.round(3)
    ext = Path(path).suffix.lower()
    if ext != ".xlsx":
        # CSV can contain only one table, so export the concise summary when available.
        df = summary_df if not summary_df.empty else detail_df
        df.to_csv(path, index=False, encoding="utf-8-sig")
        return

    info_rows = []
    if config is not None:
        info_rows = [
            {"项目": "物料编码", "内容": config.material_code},
            {"项目": "配方名称", "内容": getattr(config, "recipe_name", "")},
            {"项目": "配方版本", "内容": getattr(config, "recipe_version", "")},
            {"项目": "配方验证状态", "内容": "已验证/正式生产" if getattr(config, "recipe_validation_status", "Draft") == "Validated" else "草稿/未验证"},
            {"项目": "结果用途", "内容": "正式判定" if getattr(config, "recipe_validation_status", "Draft") == "Validated" else "试测/未验证配方，不作正式判定"},
            {"项目": "工序", "内容": config.process_name},
            {"项目": "测量设备型号", "内容": config.equipment_model},
            {"项目": "设备校准日期", "内容": config.calibration_date},
            {"项目": "操作人员", "内容": config.operator_name},
            {"项目": "测量模式", "内容": _mode_cn(config.mode)},
            {"项目": "工作方式", "内容": "自动识别测量" if getattr(config, "workflow_mode", "Manual") == "Auto" else "手动 ROI 测量"},
            {"项目": "识别质量门槛", "内容": quality_profile_display(config)},
            {"项目": "最低置信度", "内容": config.confidence_min},
            {"项目": "最低覆盖率", "内容": config.production_min_coverage},
            {"项目": "最大异常点比例", "内容": config.production_max_rejected_ratio},
            {"项目": "最大残差(μm)", "内容": config.production_max_residual_um},
            {"项目": "最大轮廓偏差(μm)", "内容": config.production_max_radial_deviation_um},
            {"项目": "自动基准 Mark", "内容": getattr(config, "auto_reference_label", "")},
            {"项目": "自动待测 Mark", "内容": getattr(config, "auto_target_label", "")},
            {"项目": "像素尺寸X(μm/px)", "内容": config.pixel_size_x_um},
            {"项目": "像素尺寸Y(μm/px)", "内容": config.pixel_size_y_um},
            {"项目": "Rx角度(μrad)", "内容": getattr(config, "rx_angle_urad", 0.0)},
            {"项目": "Ry角度(μrad)", "内容": getattr(config, "ry_angle_urad", 0.0)},
            {"项目": "物料厚度(mm)", "内容": getattr(config, "material_thickness_mm", 0.0)},
            {"项目": "角度补偿公式", "内容": "ΔX=原始ΔX+厚度×Ry/1000；ΔY=原始ΔY-厚度×Rx/1000"},
            {"项目": "椭圆直径定义", "内容": "(物理长轴+物理短轴)/2"},
            {"项目": "椭圆圆度定义", "内容": "(物理长轴-物理短轴)/2；非 ISO 最小区域圆度"},
            {"项目": "Rz分布方向", "内容": config.rz_layout},
            {"项目": "Rz单位", "内容": "μrad"},
            {"项目": "Mark间距L(μm)", "内容": config.rz_distance_l_um},
        ]
    if traceability_info:
        info_rows.extend([
            {"项目": "测量记录编号", "内容": traceability_info.get("measurement_id", "")},
            {"项目": "运行模式", "内容": traceability_info.get("operation_mode", "")},
            {"项目": "配方SHA256", "内容": traceability_info.get("recipe_hash", "")},
            {"项目": "追溯档案路径", "内容": traceability_info.get("archive_path", "")},
        ])
    info_df = pd.DataFrame(info_rows).round(3)

    with TemporaryDirectory() as tmp_dir:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            info_df.to_excel(writer, index=False, sheet_name="基础信息")
            summary_df.to_excel(writer, index=False, sheet_name="结果汇总")
            if not repeatability_df.empty:
                repeatability_df.to_excel(writer, index=False, sheet_name="多次测量结果")
            detail_df.to_excel(writer, index=False, sheet_name="识别明细")
            if not geometry_df.empty:
                geometry_df.to_excel(writer, index=False, sheet_name="尺寸结果")

            sheet_names = ["基础信息", "结果汇总"]
            if not repeatability_df.empty:
                sheet_names.append("多次测量结果")
            sheet_names.append("识别明细")
            if not geometry_df.empty:
                sheet_names.append("尺寸结果")
            for sheet_name in sheet_names:
                ws = writer.book[sheet_name]
                _style_sheet(ws, fail_columns=["结果", "判定", "提示"])

            if mark_images:
                ws = writer.book.create_sheet("图像工作区")
                ws["A1"] = "标记"
                ws["B1"] = "层"
                ws["C1"] = "完整图像与识别叠加"
                ws["D1"] = "说明"
                _style_sheet(ws)
                for idx, item in enumerate(mark_images, start=2):
                    ws.cell(idx, 1, item.get("mark_id", ""))
                    ws.cell(idx, 2, item.get("layer", ""))
                    ws.cell(idx, 4, item.get("note", ""))
                    image_path = item.get("path")
                    if image_path and Path(image_path).exists():
                        img = XlsxImage(str(image_path))
                        img.width, img.height = resize_dimensions_preserving_aspect(img.width, img.height, 720.0, 480.0)
                        ws.add_image(img, f"C{idx}")
                        ws.row_dimensions[idx].height = max(60, img.height * 0.75 + 10)
                ws.column_dimensions["C"].width = 96
