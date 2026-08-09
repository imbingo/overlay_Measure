from __future__ import annotations

import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from overlay_measure.batch_results import compact_detection
from overlay_measure.candidate_ordering import (
    assign_spatial_candidate_ids,
    candidate_display_label,
    resolve_preferred_candidate,
)
from overlay_measure.image_loader import load_image
from overlay_measure.measurement_engine import run_measurement_job
from overlay_measure.measurement_service import _fit_to_detection
from overlay_measure.measurement_units import ellipse_metrics_um
from overlay_measure.circle_ellipse_fitter import FitResult
from overlay_measure.models import DetectionResult, MarkRecipe, MeasurementConfig, Roi
from overlay_measure.recipe_manager import load_recipe
from overlay_measure.result_exporter import build_detection_rows
from overlay_measure.ui_main import MainWindow


SAMPLE_DIR = Path(__file__).resolve().parents[1] / "sample_data"


def _detection(x: float, y: float, layer: str = "upper", diameter: float = 20.0) -> DetectionResult:
    return DetectionResult(
        "Mark1", layer, x, y, x * 0.1, y * 0.1, diameter, diameter * 0.1,
        0.1, 0.01, 64, 0.95, "ProductionCircle", shape_params={"quality_status": "Valid"},
    )


def test_ellipse_metrics_use_physical_axes_and_preserve_major_order():
    config = MeasurementConfig(pixel_size_x_um=0.2, pixel_size_y_um=0.1)
    metrics = ellipse_metrics_um({"major_px": 10.0, "minor_px": 8.0, "angle_deg": 0.0}, config)
    assert metrics == pytest.approx({
        "ellipse_major_um": 2.0,
        "ellipse_minor_um": 0.8,
        "ellipse_diameter_um": 1.4,
        "ellipse_roundness_um": 0.6,
    })

    rotated = ellipse_metrics_um({"major_px": 10.0, "minor_px": 8.0, "angle_deg": 90.0}, config)
    assert rotated == pytest.approx({
        "ellipse_major_um": 1.6,
        "ellipse_minor_um": 1.0,
        "ellipse_diameter_um": 1.3,
        "ellipse_roundness_um": 0.3,
    })

    fit = FitResult(10, 20, 9, 0.1, "Ellipse", 0.95, {"major_px": 10.0, "minor_px": 8.0, "angle_deg": 0.0})
    detected = _fit_to_detection("Mark1", "upper", fit, [], config, Roi(0, 0, 30, 30))
    assert detected.ellipse_major_um == pytest.approx(2.0)
    assert detected.ellipse_minor_um == pytest.approx(0.8)
    assert detected.ellipse_diameter_um == pytest.approx(1.4)
    assert detected.ellipse_roundness_um == pytest.approx(0.6)


def test_export_roundness_is_ellipse_only():
    config = MeasurementConfig(pixel_size_x_um=0.2, pixel_size_y_um=0.1)
    ellipse = _detection(10, 20)
    ellipse.fitting_mode = "Ellipse"
    ellipse.shape_params.update({"major_px": 10.0, "minor_px": 8.0, "angle_deg": 0.0})
    circle = _detection(30, 20)
    rows = build_detection_rows(
        {"椭圆": {"upper": ellipse}, "圆": {"upper": circle}}, {}, config, run_index=2
    )
    assert rows[0]["run_index"] == 2
    assert rows[0]["ellipse_diameter_um"] == pytest.approx(1.4)
    assert rows[0]["ellipse_roundness_um"] == pytest.approx(0.6)
    assert rows[1]["ellipse_roundness_um"] is None


def test_candidates_are_row_major_but_keep_detection_rank_order():
    detections = [
        _detection(200, 100, diameter=30),
        _detection(100, 200, diameter=29),
        _detection(100, 100, diameter=28),
        _detection(300, 200, diameter=27),
        _detection(300, 100, diameter=26),
        _detection(200, 200, diameter=25),
    ]
    entries = assign_spatial_candidate_ids(detections, dual_image=False)
    assert [item.shape_params["detection_rank"] for _, item in entries] == list(range(6))
    labels_by_position = {
        (item.center_x_px, item.center_y_px): candidate_display_label(key, item)
        for key, item in entries
    }
    assert labels_by_position == {
        (100, 100): "1", (200, 100): "2", (300, 100): "3",
        (100, 200): "4", (200, 200): "5", (300, 200): "6",
    }
    mapping = {key: {item.layer: item} for key, item in entries}
    assert resolve_preferred_candidate("a", mapping) == entries[0][0]
    assert resolve_preferred_candidate("Mark1-b", mapping) == entries[1][0]


def test_dual_image_numbering_restarts_per_layer():
    entries = assign_spatial_candidate_ids(
        [_detection(30, 20, "lower"), _detection(20, 20, "upper"), _detection(10, 20, "lower")],
        dual_image=True,
    )
    displays = {item.layer + str(item.center_x_px): candidate_display_label(key, item) for key, item in entries}
    assert displays == {"lower30": "下-2", "upper20": "上-1", "lower10": "下-1"}


def test_compact_detection_removes_large_caliper_payload():
    detection = _detection(10, 10)
    detection.edge_points = [(float(i), 1.0) for i in range(1000)]
    detection.edge_gradients = [1.0] * 1000
    detection.shape_params.update({
        "caliper_windows": [{"center_x": float(i)} for i in range(1000)],
        "candidate_contour_points": [(float(i), 2.0) for i in range(1000)],
    })
    compact = compact_detection(detection)
    assert len(compact.edge_points) == 256
    assert len(compact.shape_params["candidate_contour_points"]) == 256
    assert compact.edge_gradients == []
    assert "caliper_windows" not in compact.shape_params


def test_batch_engine_retains_each_detection_snapshot():
    config, params, marks = load_recipe(str(SAMPLE_DIR / "demo_recipe.json"))
    config.workflow_mode = "Manual"
    config.recipe_validation_status = "Validated"
    upper = load_image(str(SAMPLE_DIR / "sample_upper.png"))
    lower = load_image(str(SAMPLE_DIR / "sample_lower.png"))
    payload = run_measurement_job(
        {
            "config": config,
            "params": params,
            "marks": {"Mark1": marks[0], "Mark2": MarkRecipe("Mark2")},
            "mark_images": {"Mark1": {"upper": None, "lower": None}, "Mark2": {"upper": None, "lower": None}},
            "batch_images": {
                "Mark1": {"upper": [upper, upper], "lower": [lower, lower]},
                "Mark2": {"upper": [], "lower": []},
            },
            "selections": {"Mark1": {}, "Mark2": {}},
            "batch": True,
        },
        lambda *args: None,
        lambda: False,
    )
    records = payload["batch_records"]["Mark1"]
    assert [record["run_index"] for record in records] == [1, 2]
    assert all(
        {detection.layer for detection in record["detections"].values()} == {"upper", "lower"}
        for record in records
    )
    assert all(record["selection"] for record in records)
    assert all(
        all(len(detection.edge_points) <= 256 for detection in record["detections"].values())
        for record in records
    )
    assert payload["overlays"]["Mark1"].delta_x_um == pytest.approx(
        records[0]["overlay"].delta_x_um
    )


def test_batch_detail_selector_switches_preview_and_supports_all():
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window._set_combo_value(window.measurement_run_mode_combo, "Batch")
    first = _detection(10, 10)
    second = _detection(20, 20)
    window.batch_run_records["Mark1"] = [
        {"run_index": 1, "workflow": "Manual", "detections": {"upper": first}, "selection": {}, "upper_file": "one.png", "lower_file": "", "error": ""},
        {"run_index": 2, "workflow": "Manual", "detections": {"upper": second}, "selection": {}, "upper_file": "two.png", "lower_file": "", "error": ""},
    ]
    window._refresh_batch_detail_selector(default_first=True)
    assert window.batch_detail_combo.currentData() == 1
    assert len(window._display_detection_entries()) == 1
    window.batch_detail_combo.setCurrentIndex(window.batch_detail_combo.findData(2))
    assert window._display_detection_entries()[0]["detection"].center_x_px == 20
    window.batch_detail_combo.setCurrentIndex(window.batch_detail_combo.findData("all"))
    assert len(window._display_detection_entries()) == 2
    assert "图像预览：第2次" in window.batch_detail_preview_label.text()
    window.close()
    app.processEvents()


def test_ui_excel_export_contains_every_batch_detection(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.config.workflow_mode = "Manual"
    window._set_combo_value(window.workflow_combo, "Manual")
    window._set_combo_value(window.measurement_run_mode_combo, "Batch")
    first = _detection(10, 10)
    second = _detection(20, 20)
    window.batch_run_records["Mark1"] = [
        {"run_index": 1, "workflow": "Manual", "detections": {"upper": first}, "selection": {}, "upper_file": "one.png", "lower_file": "", "overlay": None, "error": ""},
        {"run_index": 2, "workflow": "Manual", "detections": {"upper": second}, "selection": {}, "upper_file": "two.png", "lower_file": "", "overlay": None, "error": ""},
        {"run_index": 3, "workflow": "Manual", "detections": {}, "selection": {}, "upper_file": "bad.png", "lower_file": "", "overlay": None, "error": "识别失败"},
    ]
    output = tmp_path / "batch.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(output), "Excel (*.xlsx)"))
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.Ok)
    window.export_result_file()

    workbook = load_workbook(output, read_only=True)
    rows = list(workbook["识别明细"].iter_rows(values_only=True))
    headers = list(rows[0])
    run_column = headers.index("测量次数")
    failure_column = headers.index("失效原因")
    assert [row[run_column] for row in rows[1:]] == [1, 2, 3]
    assert rows[-1][failure_column] == "识别失败"
    info = {row[0]: row[1] for row in workbook["基础信息"].iter_rows(min_row=2, values_only=True)}
    assert info["椭圆圆度定义"].startswith("(物理长轴-物理短轴)/2")
    workbook.close()
    window.close()
    app.processEvents()
